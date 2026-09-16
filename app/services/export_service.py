"""Génération du classeur Excel récapitulatif des charges d'enseignement
du département — un fichier structuré en 15 colonnes, une ligne par
(enseignant, UE), trié par ordre alphabétique du nom de l'enseignant.

Quand un enseignant intervient sur plusieurs UE, les colonnes qui ne
dépendent pas de l'UE (nom, grade, total heure dû, total heure par
enseignant, heure(s) complémentaire(s)) sont fusionnées sur l'ensemble de
ses lignes plutôt que répétées."""

from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.intervention_charge import compute_intervention_charge
from app.repositories.interfaces import IRepository

EN_TETES = [
    "N°",
    "Nom et prénom",
    "Grade",
    "Total heure dû",
    "UE",
    "CM (h)",
    "Nombre de groupes TD",
    "Nombre d'heures TD",
    "Nombre de groupes TP",
    "Nombre d'heures TP",
    "Nombre de groupes TPE",
    "Nombre d'heures TPE",
    "Total d'heure par UE",
    "Total heure par enseignant",
    "Heure(s) complémentaire(s)",
]

LARGEURS = [5, 24, 18, 14, 32, 9, 11, 11, 11, 11, 12, 12, 13, 15, 16]

# Colonnes (1-indexées) qui ne dépendent pas de l'UE : à fusionner sur tout
# le bloc d'un enseignant quand il a plusieurs UE.
COLONNES_A_FUSIONNER = [2, 3, 4, 14, 15]


class ExportService:
    """Construit le classeur Excel à partir des mêmes données que le
    tableau de bord (Single Responsibility : cette classe ne fait que la
    mise en forme du fichier, le calcul métier reste dans le domaine)."""

    def __init__(self, enseignant_repository: IRepository, ue_repository: IRepository):
        self._enseignant_repo = enseignant_repository
        self._ue_repo = ue_repository

    def _lignes_par_ue(self, enseignant):
        par_ue = {}
        for interv in enseignant.interventions:
            ue = self._ue_repo.get_by_id(interv.ue_id)
            if not ue:
                continue
            entry = par_ue.setdefault(ue.id, {
                "ue_label": f"{ue.code} — {ue.intitule}",
                "cm": 0.0, "td_g": 0, "td_h": 0.0,
                "tp_g": 0, "tp_h": 0.0, "tpe_g": 0, "tpe_h": 0.0,
            })
            charge = compute_intervention_charge(interv, ue)
            if interv.type_enseignement == "CM":
                entry["cm"] += charge
            elif interv.type_enseignement == "TD":
                entry["td_g"] += 1
                entry["td_h"] += charge
            elif interv.type_enseignement == "TP":
                entry["tp_g"] += 1
                entry["tp_h"] += charge
            else:
                entry["tpe_g"] += 1
                entry["tpe_h"] += charge
        return sorted(par_ue.values(), key=lambda l: l["ue_label"])

    def build_workbook(self) -> BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Charges enseignants"

        ws.append(EN_TETES)
        entete_fill = PatternFill(start_color="16232F", end_color="16232F", fill_type="solid")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = entete_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

        enseignants = sorted(
            self._enseignant_repo.get_all(),
            key=lambda e: (e.nom.strip().upper(), e.prenom.strip().upper()),
        )

        centre = Alignment(horizontal="center", vertical="center")

        for ordre, e in enumerate(enseignants, start=1):
            lignes = self._lignes_par_ue(e)
            total_general = sum(l["cm"] + l["td_h"] + l["tp_h"] + l["tpe_h"] for l in lignes)
            heure_complementaire = round(total_general - e.volume_horaire_reference, 2)
            nom_complet = f"{e.prenom} {e.nom}"

            ligne_debut = ws.max_row + 1

            if not lignes:
                ws.append([
                    ordre, nom_complet, e.grade or "", e.volume_horaire_reference,
                    "—", 0, 0, 0, 0, 0, 0, 0, 0,
                    round(total_general, 2), heure_complementaire,
                ])
                continue

            for l in lignes:
                total_ue = round(l["cm"] + l["td_h"] + l["tp_h"] + l["tpe_h"], 2)
                ws.append([
                    ordre, nom_complet, e.grade or "", e.volume_horaire_reference,
                    l["ue_label"], round(l["cm"], 2),
                    l["td_g"], round(l["td_h"], 2),
                    l["tp_g"], round(l["tp_h"], 2),
                    l["tpe_g"], round(l["tpe_h"], 2),
                    total_ue, round(total_general, 2), heure_complementaire,
                ])

            ligne_fin = ws.max_row
            if ligne_fin > ligne_debut:
                # Plusieurs UE pour cet enseignant : fusionner les colonnes
                # qui ne dépendent pas de l'UE plutôt que de les répéter.
                for col in COLONNES_A_FUSIONNER:
                    ws.merge_cells(start_row=ligne_debut, start_column=col, end_row=ligne_fin, end_column=col)
                    ws.cell(row=ligne_debut, column=col).alignment = centre

        for i, largeur in enumerate(LARGEURS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = largeur

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(EN_TETES))}{ws.max_row}"

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
